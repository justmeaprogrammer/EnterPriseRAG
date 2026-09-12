from pptx import Presentation

def clean_text(text:str)->str:
    lines = (line.strip() for line in text.splitlines())
    return "\n".join(line for line in lines if line)

def parse_ppt(file_path:str)->str:
    prs = Presentation(file_path)
    text_parts=[]
    
    for slide_num,slide in enumerate(prs.slides,start=1):
        text_parts.append(f"--------Slide:{slide_num}---------")
        
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    para_text="".join(run.text for run in para.runs)
                    if para_text:
                        text_parts.append(para_text)
                        
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        text_parts.append(cell.text)
        if slide.has_notes_slide:
            notes_text = slide.notes_slide.notes_text_frame.text
            if notes_text.strip():
                text_parts.append(f"[Notes: {notes_text}]")
                
    text = "\n".join(text_parts)
    return clean_text(text=text)


