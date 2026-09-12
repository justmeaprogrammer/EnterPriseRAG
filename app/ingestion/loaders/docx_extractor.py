from docx import Document


def clean_lines(text: str) -> str:
    """Strip each line, drop blanks."""
    lines = (line.strip() for line in text.splitlines())
    return '\n'.join(line for line in lines if line)


def parse_docx(file_path: str) -> str:
    doc = Document(file_path)
    text_parts = []

    # Regular paragraphs
    for para in doc.paragraphs:
        text_parts.append(para.text)

    # Tables (often missed if you only loop paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text_parts.append(cell.text)

    text = "\n".join(text_parts)
    return clean_lines(text)